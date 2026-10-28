#!/usr/bin/env python3
"""Build the baseline-vs-optimized spreadsheet from results/<design>/best_config.json.

Every number written here is read from a file on disk produced by a real run;
nothing is typed in by hand. Run after run.py:

    ./.venv/bin/python make_xlsx.py [-o out.xlsx]
"""

from __future__ import annotations

import argparse
import json
import math
import pathlib
import subprocess

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent
RESULTS = ROOT / "results"

ORDER = ["gcd", "aes", "jpeg", "dynamic_node", "bp_fe_top"]

HDR = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="D9E2F3")
GOOD = PatternFill("solid", fgColor="C6EFCE")
BAD = PatternFill("solid", fgColor="FFC7CE")
WARN = PatternFill("solid", fgColor="FFEB9C")
WHITE = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def enrich(rec: dict) -> dict:
    """Fill in fields the best_config.json 'best' block does not carry.

    best_config.json records the winner's score/delay/area/cells, but not its
    flop count or wall time; both live in trajectory.csv. Pull them from the
    row whose label matches, so the spreadsheet can show them without a rerun.
    """
    best = rec.get("best")
    if not best:
        return rec
    traj = RESULTS / rec["design"] / "trajectory.csv"
    if not traj.exists():
        return rec
    import csv as _csv
    with traj.open() as fh:
        for row in _csv.DictReader(fh):
            if row.get("label") != best.get("label"):
                continue
            if row.get("score") and best.get("score") is not None:
                try:
                    if abs(float(row["score"]) - float(best["score"])) > 1e-9:
                        continue
                except ValueError:
                    continue
            for src, dst, cast in (("seq_cells", "seq_cells", int),
                                   ("wall_s", "wall_time_s", float)):
                v = (row.get(src) or "").strip()
                if v:
                    try:
                        best[dst] = cast(v)
                    except ValueError:
                        pass
            break
    return rec


def load() -> list[dict]:
    rows = []
    for d in ORDER:
        f = RESULTS / d / "best_config.json"
        if f.exists():
            try:
                rows.append(enrich(json.loads(f.read_text())))
            except json.JSONDecodeError:
                pass
    # include any design not in ORDER
    for p in sorted(RESULTS.iterdir()) if RESULTS.exists() else []:
        f = p / "best_config.json"
        if p.is_dir() and f.exists() and p.name not in ORDER:
            try:
                rows.append(enrich(json.loads(f.read_text())))
            except json.JSONDecodeError:
                pass
    return rows


def pct(new: float | None, old: float | None) -> float | None:
    if not new or not old:
        return None
    return 100.0 * (new / old - 1.0)


def autosize(ws, widths: dict[int, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def sheet_summary(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet("Summary", 0)
    cols = [
        "Design", "Top module", "Clock (ns)",
        "Base cells", "Opt cells", "Cells %",
        "Base area (um2)", "Opt area (um2)", "Area %",
        "Base delay (ps)", "Opt delay (ps)", "Delay %",
        "Base flops", "Opt flops",
        "Base depth", "Opt depth",
        "Base score", "Opt score", "Score gain %",
        "w_delay", "w_area", "Runs", "Verdict",
    ]
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HDR
        cell.font = WHITE
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        cell.border = BOX
    ws.freeze_panes = "B2"

    for r in rows:
        b = r.get("baseline") or {}
        o = r.get("best") or {}
        if not o:
            ws.append([r.get("design"), r.get("top", ""), r.get("clk_ns"),
                       b.get("cell_count"), None, None,
                       b.get("chip_area"), None, None,
                       b.get("abc_delay_ps"), None, None,
                       b.get("seq_cells"), None,
                       b.get("logic_depth"), None,
                       1.0, None, None,
                       None, None, r.get("n_synth_runs"),
                       r.get("status", "NO RESULT")])
            continue
        w = r.get("weights") or {}
        ws.append([
            r.get("design"), r.get("top", ""), r.get("clk_ns"),
            b.get("cell_count"), o.get("cells"), pct(o.get("cells"), b.get("cell_count")),
            b.get("chip_area"), o.get("area_um2"), pct(o.get("area_um2"), b.get("chip_area")),
            b.get("abc_delay_ps"), o.get("delay_ps"), pct(o.get("delay_ps"), b.get("abc_delay_ps")),
            b.get("seq_cells"), o.get("seq_cells", b.get("seq_cells")),
            b.get("logic_depth"), o.get("logic_depth"),
            1.0, o.get("score"),
            (100.0 * (1.0 - o["score"])) if o.get("score") else None,
            w.get("delay"), w.get("area"),
            r.get("n_synth_runs"), (r.get("verifier") or {}).get("verdict"),
        ])

    last = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=last, max_col=len(cols)):
        for cell in row:
            cell.border = BOX
    for rr in range(2, last + 1):
        for cc, fmt in ((3, "0.00"), (6, "+0.0;-0.0"), (7, "#,##0.00"), (8, "#,##0.00"),
                        (9, "+0.0;-0.0"), (10, "#,##0.00"), (11, "#,##0.00"),
                        (12, "+0.0;-0.0"), (17, "0.0000"), (18, "0.0000"),
                        (19, "+0.00;-0.00"), (20, "0.00"), (21, "0.00")):
            ws.cell(row=rr, column=cc).number_format = fmt
        for cc in (4, 5, 13, 14, 15, 16, 22):
            ws.cell(row=rr, column=cc).number_format = "#,##0"
        # colour the objective score: < 1 is an improvement
        sc = ws.cell(row=rr, column=18).value
        if isinstance(sc, (int, float)):
            ws.cell(row=rr, column=18).fill = GOOD if sc < 1.0 else BAD
            ws.cell(row=rr, column=18).font = BOLD
        for cc in (9, 12):  # area % / delay %
            v = ws.cell(row=rr, column=cc).value
            if isinstance(v, (int, float)):
                ws.cell(row=rr, column=cc).fill = GOOD if v < 0 else WARN

    # aggregate row
    scores = [r["best"]["score"] for r in rows if r.get("best") and r["best"].get("score")]
    if scores:
        geo = math.exp(sum(math.log(s) for s in scores) / len(scores))
        ws.append([])
        ws.append(["Geometric mean score", "", "", "", "", "", "", "", "", "", "", "",
                   "", "", "", "", 1.0, round(geo, 4),
                   round(100 * (1 - geo), 2)])
        gr = ws.max_row
        ws.cell(row=gr, column=1).font = BOLD
        ws.cell(row=gr, column=18).font = BOLD
        ws.cell(row=gr, column=18).number_format = "0.0000"
        ws.cell(row=gr, column=18).fill = GOOD if geo < 1 else BAD
        ws.append(["Worst-case score", "", "", "", "", "", "", "", "", "", "", "",
                   "", "", "", "", 1.0, round(max(scores), 4)])
        wr = ws.max_row
        ws.cell(row=wr, column=1).font = BOLD
        ws.cell(row=wr, column=18).number_format = "0.0000"
        ws.append([f"Improved over baseline: {sum(1 for s in scores if s < 1)}/{len(scores)}"])
        ws.cell(row=ws.max_row, column=1).font = BOLD

    autosize(ws, {1: 14, 2: 22, 3: 10, 4: 11, 5: 11, 6: 9, 7: 16, 8: 16, 9: 9,
                  10: 15, 11: 15, 12: 9, 13: 11, 14: 11, 15: 11, 16: 11,
                  17: 11, 18: 11, 19: 12, 20: 9, 21: 9, 22: 8, 23: 20})


def sheet_metric_view(wb: Workbook, rows: list[dict]) -> None:
    """The requested metric-per-row layout, baseline vs optimized per design."""
    ws = wb.create_sheet("Baseline vs Optimized")
    ws.append(["Metric"] + sum(([f"{r['design']} base", f"{r['design']} opt"] for r in rows), []))
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HDR
        cell.font = WHITE
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = BOX

    def emit(name: str, bkey: str, okey: str, fmt: str = "#,##0") -> None:
        vals: list = [name]
        for r in rows:
            b = r.get("baseline") or {}
            o = r.get("best") or {}
            vals += [b.get(bkey), o.get(okey)]
        ws.append(vals)
        rr = ws.max_row
        ws.cell(row=rr, column=1).font = BOLD
        ws.cell(row=rr, column=1).fill = SUB
        for cc in range(2, ws.max_column + 1):
            ws.cell(row=rr, column=cc).number_format = fmt

    emit("Cells", "cell_count", "cells")
    emit("Chip area (um2)", "chip_area", "area_um2", "#,##0.00")
    emit("Flops (sequential cells)", "seq_cells", "seq_cells")
    emit("Combinational logic depth", "logic_depth", "logic_depth")
    emit("Critical-path delay (ps)", "abc_delay_ps", "delay_ps", "#,##0.00")
    emit("Synthesis runtime (s)", "wall_time_s", "wall_time_s", "0.00")

    # objective score row
    vals: list = ["Objective score"]
    for r in rows:
        vals += [1.0, (r.get("best") or {}).get("score")]
    ws.append(vals)
    rr = ws.max_row
    ws.cell(row=rr, column=1).font = BOLD
    ws.cell(row=rr, column=1).fill = SUB
    for cc in range(2, ws.max_column + 1):
        ws.cell(row=rr, column=cc).number_format = "0.0000"
        v = ws.cell(row=rr, column=cc).value
        if isinstance(v, (int, float)) and cc % 2 == 1:  # the "opt" columns
            ws.cell(row=rr, column=cc).fill = GOOD if v < 1.0 else BAD
            ws.cell(row=rr, column=cc).font = BOLD

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.border = BOX
    ws.freeze_panes = "B2"
    autosize(ws, {1: 28, **{c: 14 for c in range(2, ws.max_column + 1)}})


def sheet_configs(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet("Winning configs")
    dims = ["flatten", "extra_opt", "abc_D_scale", "resyn_loops", "use_dch",
            "abc_flow", "mapper", "post", "adder_map"]
    ws.append(["Design", "Label", "Score"] + dims + ["Profile", "Rules retrieved"])
    for c in range(1, ws.max_column + 1):
        ws.cell(row=1, column=c).fill = HDR
        ws.cell(row=1, column=c).font = WHITE
        ws.cell(row=1, column=c).alignment = Alignment(horizontal="center", wrap_text=True)
    for r in rows:
        o = r.get("best") or {}
        cfg = o.get("config") or {}
        prof = r.get("profile") or {}
        ws.append([
            r.get("design"), o.get("label"), o.get("score"),
            *[cfg.get(d) for d in dims],
            f"{prof.get('size_class')}/{prof.get('dominant_style')}",
            ", ".join(r.get("retrieved_rules") or []),
        ])
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.border = BOX
    for rr in range(2, ws.max_row + 1):
        ws.cell(row=rr, column=3).number_format = "0.0000"
    ws.freeze_panes = "B2"
    autosize(ws, {1: 14, 2: 12, 3: 10, 4: 9, 5: 30, 6: 12, 7: 12, 8: 10,
                  9: 11, 10: 12, 11: 22, 12: 10, 13: 20, 14: 18})


def sheet_crosscheck(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet("Baseline cross-check")
    ws.append(["This sheet verifies that the in-harness baseline really is the "
               "same configuration as the recorded resyn2 sweep."])
    ws.cell(row=1, column=1).font = BOLD
    ws.append(["The resyn2 sweep recorded no delay, so delay is measured here; "
               "cells / area / flops / depth are compared against the CSV."])
    ws.append([])
    ws.append(["Design", "CSV cells", "Measured cells", "CSV area", "Measured area",
               "Area delta %", "CSV flops", "Measured flops", "CSV depth",
               "Measured depth", "Matches CSV?"])
    hr = ws.max_row
    for c in range(1, 12):
        ws.cell(row=hr, column=c).fill = HDR
        ws.cell(row=hr, column=c).font = WHITE
        ws.cell(row=hr, column=c).alignment = Alignment(horizontal="center", wrap_text=True)
    for r in rows:
        x = r.get("baseline_csv_crosscheck") or {}
        if not x:
            ws.append([r.get("design"), "not compared"])
            continue
        ws.append([r.get("design"), x.get("csv_cells"), x.get("measured_cells"),
                   x.get("csv_area"), x.get("measured_area"), x.get("area_delta_pct"),
                   x.get("csv_flops"), x.get("measured_flops"),
                   x.get("csv_logic_depth"), x.get("measured_logic_depth"),
                   "YES" if x.get("matches_csv") else "NO"])
        rr = ws.max_row
        ws.cell(row=rr, column=11).fill = GOOD if x.get("matches_csv") else WARN
        ws.cell(row=rr, column=11).font = BOLD
        for cc in (4, 5):
            ws.cell(row=rr, column=cc).number_format = "#,##0.00"
        ws.cell(row=rr, column=6).number_format = "0.00"
    autosize(ws, {1: 14, 2: 11, 3: 14, 4: 13, 5: 14, 6: 12, 7: 11, 8: 14,
                  9: 11, 10: 14, 11: 13})


def sheet_notes(wb: Workbook, rows: list[dict], usage: dict | None) -> None:
    ws = wb.create_sheet("Method & caveats")
    try:
        yv = subprocess.run(["yosys", "-V"], capture_output=True, text=True).stdout.strip()
    except OSError:
        yv = "unknown"
    lines = [
        ("Objective", ""),
        ("", "score = w_delay * (delay/delay_base) + w_area * (area/area_base); lower is better."),
        ("", "Baseline scores exactly 1.000 by construction. Weights are per-design, set by the profiler."),
        ("", "A config is rejected if a run fails/times out, check -assert fails, the port list or flop"),
        ("", "count changes, or a non-target metric degrades past the leeway threshold."),
        ("Baseline", ""),
        ("", "The validated resyn2 configuration from results/synth_resyn2_summary.csv, RE-MEASURED"),
        ("", "in this harness so that a delay number exists (the original sweep recorded none)."),
        ("", "See the 'Baseline cross-check' sheet for agreement with the recorded CSV."),
        ("Tooling", ""),
        ("", f"Yosys: {yv}"),
        ("", "Platform: nangate45, NangateOpenCellLibrary_typical.lib"),
        ("Caveats", ""),
        ("", "1. SYNTHESIS ONLY. Nothing is placed or routed; these gains need not survive P&R."),
        ("", "2. Delay is ABC's `stime -p` over liberty arcs at WireLoad=none: no SDC, no clock"),
        ("", "   definitions, no slack, no wire load, single corner. It ranks configs against each"),
        ("", "   other on one design reliably, which is what the optimizer needs, but it is NOT a"),
        ("", "   timing-signoff number. Yosys's own `sta` pass aborts on liberty-mapped netlists in"),
        ("", "   0.33, so there is no second timing engine in the loop."),
        ("", "3. 'Combinational logic depth' is `ltp -noff` measured AFTER deleting all sequential"),
        ("", "   cells. Raw `ltp -noff` on a mapped netlist walks through the library flops, hits"),
        ("", "   sequential loops and reports a meaningless number."),
        ("", "4. Retiming is pinned off: it changes the sequential boundary, which the flop-count"),
        ("", "   gate cannot verify."),
        ("", "5. DONT_USE_CELLS is not honoured - `abc` in Yosys 0.33 has no -dont_use. Affects"),
        ("", "   baseline and candidates equally, so ratios hold, but absolute areas differ from ORFS."),
        ("", "6. The SAT miter does not converge on these mapped netlists (all outputs come back"),
        ("", "   'unproven', which is NOT a counterexample). Structural checks always run; the"),
        ("", "   verifier verdict 'accept_unverified' reflects exactly that."),
    ]
    for a, b in lines:
        ws.append([a, b])
        if a:
            ws.cell(row=ws.max_row, column=1).font = BOLD
            ws.cell(row=ws.max_row, column=1).fill = SUB
    if usage:
        ws.append([])
        ws.append(["LLM overhead", f"{usage.get('llm_calls')} calls, "
                                   f"est. ${usage.get('est_cost_usd')}"])
        ws.cell(row=ws.max_row, column=1).font = BOLD
    autosize(ws, {1: 16, 2: 100})


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--out", default=str(ROOT / "results" / "yosys_agent_results.xlsx"))
    args = ap.parse_args()

    rows = load()
    if not rows:
        print("no results found - run run.py first")
        return 1

    usage = None
    uf = RESULTS / "llm_usage.json"
    if uf.exists():
        try:
            usage = json.loads(uf.read_text())
        except json.JSONDecodeError:
            pass
    if usage is None:
        # Fall back to the "llm usage: {...}" line a run prints, so a sweep
        # started before llm_usage.json existed still reports its overhead.
        import re
        for logname in ("sweep5.log", "gcd_run.log"):
            lf = ROOT / logname
            if not lf.exists():
                continue
            m = None
            for m in re.finditer(r"llm usage: (\{.*\})", lf.read_text(errors="ignore")):
                pass
            if m:
                try:
                    usage = json.loads(m.group(1))
                    break
                except json.JSONDecodeError:
                    pass

    wb = Workbook()
    wb.remove(wb.active)
    sheet_summary(wb, rows)
    sheet_metric_view(wb, rows)
    sheet_configs(wb, rows)
    sheet_crosscheck(wb, rows)
    sheet_notes(wb, rows, usage)

    out = pathlib.Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}  ({len(rows)} designs)")
    for r in rows:
        o = r.get("best") or {}
        print(f"  {r['design']:<14} score={o.get('score', 'n/a')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
